# -*- coding: utf-8 -*-
# Geometry + content QA for the generated deck (no renderer available on this
# machine, so overflow is estimated: CJK glyph ~1.0em, latin/digit ~0.52em).
import sys
import math
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
SLIDE_W, SLIDE_H = 13.333, 7.5


def em_w(ch, size):
    if ord(ch) > 0x2E7F:  # CJK & fullwidth
        return size * 1.0
    if ch == " ":
        return size * 0.3
    return size * 0.52


def para_text(p):
    return "".join(r.text for r in p.runs)


def est_height(tf, w_in):
    w_pts = w_in * 72
    total = 0.0
    for p in tf.paragraphs:
        txt = para_text(p)
        size = 13
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.12
        if not txt:
            total += size * 1.2
            continue
        indent = 0.25 * 72 if (p._pPr is not None and p._pPr.get("marL")) else 0
        line_w = max(w_pts - indent, 30)
        text_w = sum(em_w(c, size) for c in txt)
        lines = max(1, math.ceil(text_w / line_w))
        total += lines * size * 1.22 * ls
        if p.space_after:
            total += p.space_after.pt
    return total / 72.0


def inches(v):
    return Emu(v).inches


prs = Presentation(sys.argv[1])
problems = []
for si, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        if sh.shape_type is None:
            continue
        try:
            x, y = inches(sh.left), inches(sh.top)
            w, h = inches(sh.width), inches(sh.height)
        except Exception:
            continue
        if x < -0.02 or y < -0.02 or x + w > SLIDE_W + 0.02 or y + h > SLIDE_H + 0.02:
            problems.append(f"s{si}: OUT-OF-BOUNDS {sh.shape_id} at ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f}")
        if sh.has_text_frame and sh.text_frame.text.strip():
            est = est_height(sh.text_frame, w)
            if est > h * 1.06 + 0.06:
                first = sh.text_frame.text.strip().splitlines()[0][:28]
                problems.append(f"s{si}: OVERFLOW? box {h:.2f}in needs ~{est:.2f}in  [{first}…]")
            boxes.append((x, y, w, h, sh.text_frame.text.strip()[:20]))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[i]
            bx, by, bw, bh, bt = boxes[j]
            ox = max(0, min(ax + aw, bx + bw) - max(ax, bx))
            oy = max(0, min(ay + ah, by + bh) - max(ay, by))
            inter = ox * oy
            small = min(aw * ah, bw * bh)
            if small > 0 and inter / small > 0.18:
                problems.append(f"s{si}: TEXT-OVERLAP [{at}…] vs [{bt}…] {inter/small:.0%}")

print(f"slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
if problems:
    print(f"{len(problems)} potential issues:")
    for pr in problems:
        print(" ", pr)
else:
    print("geometry: clean")
