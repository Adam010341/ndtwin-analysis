# -*- coding: utf-8 -*-
# NDTwin progress-report deck, first draft (python-pptx port).
# Source of truth: /home/adam/Desktop/NDTwin Slide material/NDTwin-slide-template.md
# Numbers re-measured 2026-08-13 at head cfbbf24 per template rule A3.1.
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

F = "Microsoft JhengHei"

C = dict(
    navy="13293D", ink="22333F", steel="2E6FA3", steelTint="E9F1F8",
    teal="0E8A7D", tealTint="E6F4F2", red="B3392F", redTint="FBEDEB",
    slate="50697F", slateTint="EDF1F4", gray="5F7180", faint="8FA0AD",
    line="D5DEE6", card="F4F7FA", white="FFFFFF", gold="B98A1C",
    goldTint="FDF6E7", darkCard="1D3A54", iceOnDark="BFD4E6",
)

SEC = dict(
    open=("開場", C["navy"]),
    f=("1 · 新增功能", C["steel"]),
    b=("2 · Baseline 缺陷", C["red"]),
    t=("3 · 使用技術", C["teal"]),
    q=("4 · 測試與文件", C["slate"]),
    l=("5 · 實機測試與 Demo", C["navy"]),
)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

page_no = 0


def rgb(h):
    return RGBColor.from_string(h)


def set_run(run, size=13, bold=False, italic=False, color=C["ink"]):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = F  # sets a:latin
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", F)


def para_bullet(p, on=True):
    pPr = p._pPr if p._pPr is not None else p.get_or_add_pPr()
    if on:
        pPr.set("marL", "228600")
        pPr.set("indent", "-228600")
        bf = pPr.makeelement(qn("a:buFont"), {"typeface": F})
        bc = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bf)
        pPr.append(bc)
    else:
        bn = pPr.makeelement(qn("a:buNone"), {})
        pPr.append(bn)


def add_text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = valign
    first = True
    for pa in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = pa.get("align", align)
        if pa.get("gap"):
            p.space_after = Pt(pa["gap"])
        p.line_spacing = pa.get("ls", 1.12)
        para_bullet(p, on=bool(pa.get("bullet")))
        r = p.add_run()
        r.text = pa["text"]
        set_run(r, size=pa.get("size", 13), bold=pa.get("bold", False),
                italic=pa.get("italic", False), color=pa.get("color", C["ink"]))
    return tb


def add_box(slide, x, y, w, h, fill=None, line_color=None, line_w=0.75,
            radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line_color:
        sp.line.color.rgb = rgb(line_color)
        sp.line.width = Pt(line_w)
        if dash:
            sp.line.dash_style = dash
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def labeled_box(slide, x, y, w, h, text, fill=C["card"], color=C["ink"],
                size=11.5, bold=False, line_color=C["line"], radius=0.08):
    add_box(slide, x, y, w, h, fill=fill, line_color=line_color, radius=radius)
    paras = [dict(text=t, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER, ls=1.05)
             for t in text.split("\n")]
    add_text(slide, x + 0.08, y, w - 0.16, h, paras, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def arrow(slide, x1, y1, x2, y2, color=C["steel"], width=1.6, dash=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = rgb(color)
    cn.line.width = Pt(width)
    if dash:
        cn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = cn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    cn.shadow.inherit = False
    return cn


def slide_base(sec_key, dark=False, chrome=True):
    global page_no
    page_no += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(C["navy"] if dark else C["white"])
    label, color = SEC[sec_key]
    if chrome:
        add_box(s, 0.55, 0.32, 2.05, 0.34, fill=(C["darkCard"] if dark else color), radius=0.5)
        add_text(s, 0.55, 0.32, 2.05, 0.34,
                 [dict(text=label, size=10.5, bold=True,
                       color=(C["iceOnDark"] if dark else C["white"]), align=PP_ALIGN.CENTER)],
                 valign=MSO_ANCHOR.MIDDLE)
        add_text(s, 0.55, 7.14, 5.5, 0.28,
                 [dict(text="NDTwin 進度報告 · 初版草稿 2026-08-13", size=8.5,
                       color=(C["iceOnDark"] if dark else C["faint"]))],
                 valign=MSO_ANCHOR.MIDDLE)
        add_text(s, 12.3, 7.14, 0.5, 0.28,
                 [dict(text=str(page_no), size=9, color=(C["iceOnDark"] if dark else C["faint"]),
                       align=PP_ALIGN.RIGHT)], valign=MSO_ANCHOR.MIDDLE)
    return s


def title(s, txt, dark=False, y=0.74, size=26, w=12.2, h=0.62):
    add_text(s, 0.55, y, w, h, [dict(text=txt, size=size, bold=True,
             color=(C["white"] if dark else C["navy"]))], valign=MSO_ANCHOR.MIDDLE)


def kicker(s, txt, dark=False, y=1.38):
    add_text(s, 0.55, y, 12.2, 0.4, [dict(text=txt, size=13, italic=True,
             color=(C["iceOnDark"] if dark else C["gray"]))], valign=MSO_ANCHOR.MIDDLE)


def row(s, y, num, head, body, color=C["red"], x=0.6, w=12.15, h=1.16,
        head_size=13, body_size=11.5):
    add_box(s, x, y + 0.03, 0.34, 0.34, fill=color, shape=MSO_SHAPE.OVAL)
    add_text(s, x, y + 0.03, 0.34, 0.34,
             [dict(text=num, size=12, bold=True, color=C["white"], align=PP_ALIGN.CENTER)],
             valign=MSO_ANCHOR.MIDDLE)
    add_text(s, x + 0.52, y, w - 0.52, h, [
        dict(text=head, size=head_size, bold=True, color=C["ink"], gap=2),
        dict(text=body, size=body_size, color=C["gray"], ls=1.08),
    ])


def stat(s, x, y, w, h, num, label, num_color=C["steel"], num_size=30, label_size=10.5):
    add_box(s, x, y, w, h, fill=C["card"], line_color=C["line"], radius=0.09)
    add_text(s, x + 0.12, y + 0.1, w - 0.24, h * 0.56,
             [dict(text=num, size=num_size, bold=True, color=num_color, align=PP_ALIGN.CENTER)],
             valign=MSO_ANCHOR.MIDDLE)
    add_text(s, x + 0.12, y + h * 0.6, w - 0.24, h * 0.36,
             [dict(text=label, size=label_size, color=C["gray"], align=PP_ALIGN.CENTER, ls=1.0)],
             valign=MSO_ANCHOR.TOP)


def placeholder(s, x, y, w, h, detail):
    add_box(s, x, y, w, h, fill=C["goldTint"], line_color=C["gold"], line_w=1,
            radius=0.09, dash=MSO_LINE_DASH_STYLE.DASH)
    add_text(s, x + 0.25, y + 0.18, w - 0.5, h - 0.36, [
        dict(text="【Adam 後續提供】", size=13, bold=True, color=C["gold"], gap=4),
        dict(text=detail, size=11, color=C["gray"], ls=1.12),
    ])


def bullets(s, items, x=0.6, y=1.9, w=12.1, h=4.9, dark=False, size=13, gap=7):
    paras = []
    for it in items:
        paras.append(dict(
            text=it["t"], bullet=not it.get("plain"),
            size=it.get("size", size), bold=it.get("b", False),
            color=it.get("c", C["iceOnDark"] if dark else C["ink"]),
            gap=it.get("gap", gap), ls=1.12,
        ))
    add_text(s, x, y, w, h, paras)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# ============================================================
# 1 — Title (dark)
s = slide_base("open", dark=True, chrome=False)
add_box(s, 0.9, 2.02, 0.14, 1.9, fill=C["teal"], radius=0.5)
add_text(s, 1.3, 2.0, 11.3, 0.75, [dict(text="NDTwin 數位孿生", size=40, bold=True, color=C["white"])], valign=MSO_ANCHOR.MIDDLE)
add_text(s, 1.3, 2.82, 11.3, 0.6, [dict(text="P4/bmv2 資料面支援，與它揭露的 baseline 缺陷", size=22, color=C["iceOnDark"])], valign=MSO_ANCHOR.MIDDLE)
add_text(s, 1.3, 3.52, 11.3, 0.4, [dict(text="實驗室進度報告", size=14, italic=True, color=C["iceOnDark"])], valign=MSO_ANCHOR.MIDDLE)
add_text(s, 1.3, 5.6, 11.3, 0.45, [dict(text="Adam　·　指導教授：【待填】　·　2026-09-03（預定）", size=14, color=C["white"])], valign=MSO_ANCHOR.MIDDLE)
notes(s, "template Page 1。題目採 template 建議句。姓名全名與指導教授姓名待 Adam 填；日期為預定報告日（約 2026-09-03）。")

# 2 — Outline
s = slide_base("open")
title(s, "大綱")
outline_rows = [
    ("1", "新增功能", "P4/bmv2 資料面完整接入：strategy pattern、P4 pipeline、proxy agent、liveness、failover、電源管理", C["steel"]),
    ("2", "修復的 baseline bug", "接手系統既有的無聲缺陷：crash、資料正確性、狀態破壞、資源類——每條附「為什麼之前沒人發現」", C["red"]),
    ("3", "使用技術", "技術棧總覽，與這次工作真正的重點：方法論", C["teal"]),
    ("4", "測試工具與文件", "五層測試架構（L0–L4）、測試資產與品質保證、文件資產", C["slate"]),
    ("5", "實機測試與 demo", "L4 差異比對、failover 端到端、壓力測試、精確度量測、demo", C["navy"]),
]
y = 1.78
for n, h, b, col in outline_rows:
    add_box(s, 0.7, y + 0.1, 0.52, 0.52, fill=col, shape=MSO_SHAPE.OVAL)
    add_text(s, 0.7, y + 0.1, 0.52, 0.52, [dict(text=n, size=17, bold=True, color=C["white"], align=PP_ALIGN.CENTER)], valign=MSO_ANCHOR.MIDDLE)
    add_text(s, 1.5, y, 11.1, 0.95, [
        dict(text=h, size=16, bold=True, gap=2),
        dict(text=b, size=11.5, color=C["gray"]),
    ], valign=MSO_ANCHOR.MIDDLE)
    y += 1.02
notes(s, "template Page 2。")

# 3 — Background
s = slide_base("open")
title(s, "背景：NDTwin 是什麼、我的任務")
bullets(s, [
    dict(t="NDTwin＝網路數位孿生：C++ kernel 為核心，7 個外圍元件（GUI、Energy-Saving、Traffic-Engineering…）透過 /ndt/ HTTP API 互動", size=14),
    dict(t="既有系統支援 OVS/Mininet ＋ Ryu 控制器（實驗室既有成果）", size=14),
    dict(t="我的任務：讓同一顆 kernel 也能驅動 P4/bmv2 資料面", b=True, size=14),
], y=1.72, h=2.1)
yb = 4.5
labeled_box(s, 0.8, yb, 2.5, 0.8, "7 個外圍應用", fill=C["slateTint"])
labeled_box(s, 3.9, yb, 2.3, 0.8, "/ndt/ HTTP API", fill=C["card"])
labeled_box(s, 6.8, yb, 2.5, 0.8, "C++ Kernel", fill=C["navy"], color=C["white"], bold=True, line_color=None)
labeled_box(s, 9.9, yb, 2.6, 0.8, "OVS/Mininet ＋ Ryu", fill=C["steelTint"])
arrow(s, 3.3, yb + 0.4, 3.9, yb + 0.4)
arrow(s, 6.2, yb + 0.4, 6.8, yb + 0.4)
arrow(s, 9.3, yb + 0.4, 9.9, yb + 0.4)
add_text(s, 9.9, yb + 0.92, 2.6, 0.35, [dict(text="＋ P4/bmv2（本次工作）", size=11.5, bold=True, color=C["teal"], align=PP_ALIGN.CENTER)])
notes(s, "template Page 3。素材：CHANGELOG L9-20；doc/2026-07-27_p4_bmv2_support_plan.md。既有系統一句話帶過（A1.1）。")

# 4 — Two main lines
s = slide_base("open")
title(s, "工作全貌：一個任務變成兩條主線")
bullets(s, [
    dict(t="開工假設：「baseline 功能正確，P4 是純新增」", size=14),
    dict(t="實測推翻：對真正的 baseline（28b8b13）驗出 11 個既有缺陷——全部無聲：不 crash、不留 log、每個 endpoint 都回 200——這正是假設看起來成立的原因", b=True, size=14),
    dict(t="所以工作是兩件事：P4 支援 ＋ 修好共用路徑。會說謊的 baseline，無法拿來驗證新的資料面", size=14),
], y=1.72, h=2.35)
stat(s, 0.8, 4.45, 3.7, 1.75, "296", "commits（2026-07-23 → 08-13，全部本人）")
stat(s, 4.85, 4.45, 3.7, 1.75, "11", "個無聲的 baseline 既有缺陷", num_color=C["red"])
stat(s, 8.9, 4.45, 3.7, 1.75, "2", "條主線：P4 支援＋共用路徑修復", num_color=C["teal"])
notes(s, "template Page 4（整份簡報的敘事鑰匙，建議講滿）。素材：CHANGELOG L376-386 ⚠ 段。296 為 2026-08-13 於 head cfbbf24 重量（git log 28b8b13..HEAD --oneline | wc -l）；單一作者已驗（git log --format=%an | sort -u 只有 Adam010341）。28b8b13 是 6f32bca 的 direct parent（已核實，不是近似值）。")

# 5 — Architecture
s = slide_base("f")
title(s, "架構：P4 怎麼接進既有系統——kernel 無感")
ky = 2.55
labeled_box(s, 0.7, 1.62, 3.3, 0.62, "7 個外圍應用 ＋ Intent Translator", fill=C["slateTint"], size=11)
arrow(s, 2.35, 2.24, 2.35, ky, color=C["gray"], width=1.2)
labeled_box(s, 0.7, ky, 3.3, 1.75, "C++ Kernel\nClassifier / FlowLinkUsageCollector\n全部 /ndt/ API\nsFlow collector UDP:6343", fill=C["navy"], color=C["white"], size=11.5, line_color=None)
labeled_box(s, 5.35, 1.95, 3.0, 0.85, "Ryu 控制器\n北向 REST :8080", fill=C["steelTint"], size=11.5)
labeled_box(s, 5.35, 3.65, 3.0, 0.85, "P4 Proxy Agent\n模仿 Ryu API :8081", fill=C["tealTint"], size=11.5, bold=True)
labeled_box(s, 9.85, 1.95, 2.7, 0.85, "OVS / Mininet", fill=C["steelTint"], size=11.5)
labeled_box(s, 9.85, 3.65, 2.7, 0.85, "bmv2 / Mininet", fill=C["tealTint"], size=11.5, bold=True)
arrow(s, 4.0, 2.85, 5.35, 2.35)
arrow(s, 4.0, 3.7, 5.35, 4.05, color=C["teal"])
arrow(s, 8.35, 2.37, 9.85, 2.37)
add_text(s, 8.42, 2.02, 1.4, 0.3, [dict(text="OpenFlow", size=9, color=C["gray"])])
arrow(s, 8.35, 4.07, 9.85, 4.07, color=C["teal"])
add_text(s, 8.32, 3.72, 1.6, 0.3, [dict(text="P4Runtime gRPC", size=9, color=C["gray"])])
arrow(s, 9.1, 3.22, 4.0, 3.22, color=C["teal"], dash=True, width=1.4)
add_text(s, 4.35, 2.9, 4.6, 0.27, [dict(text="sFlow v5 回傳 → kernel UDP:6343（OVS 原生／proxy 合成）", size=9, color=C["teal"])])
add_box(s, 0.7, 4.62, 3.3, 0.5, fill=C["tealTint"], line_color=C["teal"], line_w=1, radius=0.1)
add_text(s, 0.7, 4.62, 3.3, 0.5, [dict(text="kernel 上層一行都不用改", size=12, bold=True, color=C["teal"], align=PP_ALIGN.CENTER)], valign=MSO_ANCHOR.MIDDLE)
bullets(s, [
    dict(t="proxy agent 模仿 Ryu 北向 API（topology、/stats/flow、destination paths）——kernel 以為自己還在跟 Ryu 講話", size=12.5, gap=5),
    dict(t="proxy 自行合成 sFlow v5，打進 kernel 既有 UDP:6343 collector——遙測路徑也零修改", size=12.5, gap=5),
    dict(t="結果：Classifier、FlowLinkUsageCollector、全部 /ndt/ API、7 個外圍應用、Intent Translator 一行都不用改", size=12.5, b=True, gap=5),
], y=5.35, h=1.7)
notes(s, "template Page 5。素材：CHANGELOG L11-15；p4_proxy/proxy_agent/SPEC.md。Port 佈局（兩輪 live＋程式碼證實）：kernel :8000、P4 proxy :8081、Ryu :8080。")

# 6 — Strategy pattern
s = slide_base("f")
title(s, "Strategy pattern 與 SwitchKind 分派")
bullets(s, [
    dict(t="IRoutingStrategy / IPowerStrategy 介面，OVS 與 P4 雙實作——路由與電源各自一條乾淨的分派線", size=14),
    dict(t="typed SwitchKind（OVS / BMV2 / HARDWARE）取代原本「對拓撲檔名做大小寫敏感的子字串比對」", size=14),
    dict(t="O(1) 查找；不再對整張圖做逐操作深拷貝", size=14),
    dict(t="未知 dpid 回錯誤，而非默默轉給 Ryu；載入時驗證拓撲同質性", size=14),
])
notes(s, "template Page 6。素材：CHANGELOG 條目 6；src/ndt_core/routing_management/、power_management/。")

# 7 — P4 pipeline
s = slide_base("f")
title(s, "P4 pipeline（ndtwin_switch.p4）")
bullets(s, [
    dict(t="ternary flow_5tuple 表（帶真正 priority）置於 ipv4_lpm 之前——flow 級規則能蓋過路由", size=14),
    dict(t="ARP / TCP / UDP / ICMP 解析；L2 表——非 IPv4 frame 不再被默默丟棄", size=14),
    dict(t="TTL guard；direct / per-port counters", size=14),
    dict(t="1/256 clone-to-CPU 取樣，供遙測合成 sFlow", size=14),
], h=4.2)
add_text(s, 0.6, 6.35, 12, 0.35, [dict(text="p4_proxy/p4_src/ndtwin_switch.p4 · 482 行 · p4_src/SPEC.md", size=10.5, color=C["faint"])])
notes(s, "template Page 7。素材：CHANGELOG 條目 7。")

# 8 — Proxy modules
s = slide_base("f")
title(s, "P4 proxy agent（Python 服務）")
kicker(s, "模組構成——每個模組對應一件 kernel 期待 Ryu 做的事")
mods = [
    ("p4_client.py", "P4Runtime gRPC client：mastership、pipeline、table 讀寫"),
    ("topology_manager.py", "拓撲管理 ＋ LLDP beacon"),
    ("ryu_topology.py / ryu_flow_stats.py", "Ryu 相容 REST（含 string-action 格式的 /stats/flow）"),
    ("kernel_notifier.py", "switch-entered / link 事件通知 kernel"),
    ("sflow_emitter.py", "sFlow v5 合成"),
    ("FastAPI / uvicorn", "服務框架"),
]
for i, (hh, bb) in enumerate(mods):
    cx = 0.7 + (i % 3) * 4.1
    cy = 2.0 + (i // 3) * 1.85
    add_box(s, cx, cy, 3.85, 1.6, fill=C["card"], line_color=C["line"], radius=0.09)
    add_text(s, cx + 0.22, cy + 0.16, 3.41, 1.3, [
        dict(text=hh, size=12.5, bold=True, color=C["steel"], gap=4),
        dict(text=bb, size=11, color=C["gray"], ls=1.1),
    ])
notes(s, "template Page 8。素材：p4_proxy/proxy_agent/ 各檔；CHANGELOG 條目 26。")

# 9 — sFlow synthesis
s = slide_base("f")
title(s, "sFlow v5 合成與跨語言驗證")
bullets(s, [
    dict(t="Python 端合成的封包，與 OVS 發出的 byte-layout 相容——kernel 的 collector 分不出來源", size=14),
    dict(t="證明方式：跨語言 round-trip——把 Python emitter 的真實輸出，餵進 kernel 實際使用的 C++ parser", b=True, size=14),
    dict(t="不是「兩邊各自測」：同一串 bytes 走完整條真實路徑，layout 錯一個 byte 測試就紅", size=14),
], h=4.2)
add_text(s, 0.6, 6.35, 12, 0.35, [dict(text="p4_proxy/proxy_agent/sflow_emitter.py ←→ tests/test_SFlowEmitterRoundtrip.cpp", size=10.5, color=C["faint"])])
notes(s, "template Page 9。素材：CHANGELOG 條目 8。")

# 10 — Sampling path
s = slide_base("f")
title(s, "遙測取樣路徑：clone session 與 packet_in")
bullets(s, [
    dict(t="PRE clone session 250 經 P4Runtime 程式化——取樣不佔 pipeline 正常轉發", size=14),
    dict(t="樣本與真 packet-in 用 packet_in 內的 reason 欄位區分", size=14),
    dict(t="踩到的坑：第三個 controller header 編譯會過、但被 P4Runtime 默默忽略——它按名字只認 packet_in / packet_out", size=14),
])
notes(s, "template Page 10（可略）。素材：CHANGELOG 條目 9。「編譯過但被默默忽略」是好的口頭故事。")

# 11 — Liveness
s = slide_base("f")
title(s, "證據式存活偵測（liveness）")
labeled_box(s, 0.7, 1.7, 2.0, 0.55, "Up", fill=C["tealTint"], color=C["teal"], bold=True, size=13)
labeled_box(s, 2.9, 1.7, 2.0, 0.55, "Down", fill=C["redTint"], color=C["red"], bold=True, size=13)
labeled_box(s, 5.1, 1.7, 2.3, 0.55, "Unknown", fill=C["slateTint"], color=C["slate"], bold=True, size=13)
add_text(s, 7.7, 1.7, 5.0, 0.55, [dict(text="「無法判斷」絕不能報成「死了」——Unknown 不動 graph", size=12.5, bold=True)], valign=MSO_ANCHOR.MIDDLE)
bullets(s, [
    dict(t="bmv2 側：GET /p4/switch_state——P4Runtime round-trip 探測 ＋ LLDP 新鮮度，取代早期的無條件 setVertexUp", size=13.5),
    dict(t="死掉的 switch 從 /v1.0/topology/switches 消失，而不是留在清單裡假裝存在——live 驗證：678 筆 10 Hz 取樣、零 up-blip", size=13.5),
    dict(t="只有明確的 False 才排除一台：「沒探過」不是死亡證據——否則 fabric 在每次啟動的頭幾秒會全黑", size=13.5),
    dict(t="單一 proxy 失聯不能把十台 switch 全標黑；OVS 側套同一套 policy（該側原始行為見第 2 節）", size=13.5),
], y=2.55, h=3.9)
notes(s, "template Page 11 ＋ commit 32afeb9 的補充。素材：CHANGELOG 條目 26/27（a8db425）、32afeb9 完整 message。誠實邊界（Q&A 防線）：32afeb9 修的是 proxy 側（違約的是呼叫端，render_switches 的 docstring 本來就寫「連不到的不該出現」）；kernel 側 updateSwitches 對列出的 dpid 仍無條件 isUp=true——是「不再餵它壞資料」，不是「它自己會判斷」。量測手法：10 Hz 取樣看 up-blip 跟著 5s/30s 輪詢節奏出現＝用 cadence 指認寫入者。")

# 12 — Failover Phase 6
s = slide_base("f")
title(s, "Failover（Phase 6 · 六項全部完成）")
bullets(s, [
    dict(t="switch-entered 通知 · link failure/recovery 通知 · 證據式 liveness · LLDP beacon · Ryu 格式 /stats/flow · destination paths", b=True, size=13.5),
    dict(t="LLDP beacon 細節：port 由拓撲導出；source MAC 不再與 host 區段衝突；修掉 dpid ≥ 256 的 crash", size=13.5),
    dict(t="配套修復：拓撲變化時重算路由——既有 Ryu 控制程式原本整支沒有 remove_edge，斷鏈之後什麼都不會發生", size=13.5),
    dict(t="端到端驗證數據 → 第 5 節（bmv2 16.63 秒自癒、復原 hitless）", size=13.5, c=C["steel"]),
])
notes(s, "template Page 12。素材：CHANGELOG L426-431、條目 21；commit ca72d22（端到端驗證）、2c81b26（重算路由；註明既有 Ryu 控制程式的缺陷）。")

# 13 — Phase 7 power
s = slide_base("f")
title(s, "Phase 7：交換器電源管理")
kicker(s, "數位孿生要能回答「關掉這台會怎樣」——就必須真的能關、能開、能重新接管")
my = 1.95
labeled_box(s, 0.7, my, 3.7, 1.05, "ndtwin-p4-power helper\n關閉整台 bmv2，保留可重啟狀態", fill=C["tealTint"], size=11.5)
labeled_box(s, 4.85, my, 3.7, 1.05, "POST /p4/readopt/{dpid}\n重建 mastership / pipeline /\nclone session / 路由", fill=C["tealTint"], size=11.5)
labeled_box(s, 9.0, my, 3.6, 1.05, "死掉的 switch 從拓撲清單消失\n不留在清單裡假裝存在", fill=C["tealTint"], size=11.5)
arrow(s, 4.4, my + 0.52, 4.85, my + 0.52, color=C["teal"])
arrow(s, 8.55, my + 0.52, 9.0, my + 0.52, color=C["teal"])
stat(s, 0.7, 3.35, 3.7, 1.5, "1.50 秒", "關機 135 秒後，power-on 到重新接管", num_color=C["teal"], num_size=26)
stat(s, 4.85, 3.35, 3.7, 1.5, "0 up-blip", "678 筆 10 Hz 取樣期間", num_color=C["teal"], num_size=26)
stat(s, 9.0, 3.35, 3.6, 1.5, "9000/9000", "其餘九台零封包遺失（早輪驗收）", num_color=C["teal"], num_size=26)
bullets(s, [
    dict(t="誠實一句：readopt 在 LLDP 重新發現鏈路之前，回報「已接管、路由尚未安裝」——約 30 秒空窗，已知且可觀測", size=12.5, gap=5),
    dict(t="修復故事：修復前關機夠久後 powerOn 必定失敗，錯誤是「Connection refused」——打在一個正常 accept 的 port 上。根因是 gRPC 全域 subchannel pool 讓全新 client 繼承 120 秒 backoff；對照實驗：帶 use_local_subchannel_pool 的 channel 0.00 秒 READY、不帶的 32.56 秒", size=12.5, gap=5),
], y=5.05, h=1.9)
notes(s, "template Page 12b ①–④。素材：doc/2026-07-27_p4_bmv2_support_plan.md Phase 7 節；doc/2026-08-11_phase7_power_mechanism_design.md；commit 949fcba（subchannel pool）、e7d564b、32afeb9、3674ddd。兩輪數據不混講：1.50 秒與 678 筆是 08-12/13 輪（A-live-runbook / W2-live-reverify）；9000/9000 與「helper 回報 stopped 前 0.3 秒停止轉送」是較早那輪 Phase 7 驗收（design doc 的 Live 驗收結果）——投影片已標「早輪驗收」。gRPC 的 channel option 拼錯會被安靜忽略——所以測試明寫完整 option 字串。")

# 14 — Robustness
s = slide_base("f")
title(s, "穩健性：單一 switch 的故障，不該變成全 fabric 的狀態遺失", size=24)
bullets(s, [
    dict(t="事件：一台 bmv2 活著但不回應 gRPC → flow-table 端點是 async def，阻塞的 gRPC 讀跑在 event loop 上 → 整個 proxy 被拖死：liveness 端點 1.9 ms → 完全無回應，圖 40/40 → 32/40", size=13, b=True),
    dict(t="指認：對活著的 proxy 取 stack dump，直接指出卡住的那一格；同一份 dump 證明 liveness prober 全程健康、threadpool 全程閒置——證據早就備好了，只是沒人能讀", size=13),
    dict(t="修法：阻塞工作移出 event loop ＋ streaming gRPC 讀加 deadline → 實測 liveness 全程 1.2–1.9 ms、圖全程 40/40", size=13),
    dict(t="同頁第二例——readopt：live 驗證抓到「對健康 switch 清表、仍回 success」→ 修成安全失敗（非 primary → 502 不碰 switch；路由全拒 → 502），拒絕與接受兩個方向都補了測試", size=13),
    dict(t="⚠ 與第 2 節「無界 curl」是同一事件的兩半：這半在自己寫的 proxy 碼、那半在 baseline 的 kernel 碼——不是重複計數", size=12.5, c=C["red"]),
], y=1.8, h=4.8)
notes(s, "template Page 12b ⑤ ＋ A2 排除表的 readopt 條（裁定放穩健性頁）。素材：1a7d815（py-spy stack dump 證據、三層修法、live 驗證 /stats/flow/5 503 in 5.004s、switch_state 全程 1.2-1.9ms）、a72a168（readopt 安全失敗）。與 bug 節「API 錯誤處理」頁 ④ 互為一體兩面——兩頁都有明說。")

# 15 — Other features
s = slide_base("f")
title(s, "其他功能（快速帶過）")
bullets(s, [
    dict(t="headless 啟動：--mode / --topology / --ai / --no-ai 取代互動式 std::cin——CI 跑得動", size=14),
    dict(t="P4 明示自己的極限：group / meter 回 501 unsupported，而非默默轉給 Ryu", size=14),
    dict(t="全 bmv2 拓撲的 identity ifIndex→port 對應：跳過不認識 bmv2 的 ovs-vsctl；惰性決定，避免與拓撲載入 race", size=14),
    dict(t="南向失敗可見化：OpResult {ok, httpStatus, message}——200 但 body 是 error 也算失敗；後續把整條鏈接通", size=14),
])
notes(s, "template Page 13。素材：CHANGELOG 條目 4、10、11、12、19；8c25dbc（OpResult 接通）。OpResult 放這頁（功能完善）不放 bug 頁——A2 裁定。")

# 16 — Section 2 opener (dark)
s = slide_base("b", dark=True)
title(s, "前提翻案：baseline 有 11 個無聲缺陷", dark=True)
bullets(s, [
    dict(t="開工假設「模擬器功能正確」——對真正的 baseline（28b8b13）驗證後：11 個缺陷本來就在", size=14.5),
    dict(t="共同特徵：全部無聲。不 crash、不留 log、每個 endpoint 都回 200——所以假設看起來成立", b=True, size=14.5),
    dict(t="其後的獨立審查又坐實了多條同樣早於接手的缺陷，本節一併列出；每條都附「為什麼之前沒人發現」", size=14.5),
], y=1.75, h=2.3, dark=True)
add_box(s, 0.7, 4.25, 11.9, 2.25, fill=C["darkCard"], radius=0.1)
add_text(s, 1.0, 4.45, 11.3, 1.9, [
    dict(text="引例：一個功能在五個層次同時是死的（歷史鏈路資料）", size=14, bold=True, color=C["white"], gap=5),
    dict(text="成員從未賦值（null）· main 建了兩個不同實例、start 的和交出去的不是同一個 · 開關寫了沒人讀 · 建不出目錄就在啟動時丟例外 · 寫入不檢查 is_open——每次執行都往死掉的 stream 附加、什麼都沒寫、一行 log 都沒有", size=12, color=C["iceOnDark"], gap=5, ls=1.15),
    dict(text="「每一層單獨都足以讓功能失效——所以修好任何一層，都不會有任何改善。」修一層的人會以為自己修錯了地方。", size=12.5, italic=True, color=C["white"]),
])
notes(s, "template Page 14 ＋ ee7233b 當開場引例（補充審查報告建議，A2 已列入）。素材：CHANGELOG L376-386、L388-410；ee7233b 完整 message（兩層已逐字對 28b8b13 驗證：ControllerAndOtherEventHandler.cpp:51/:68 成員未賦值；main.cpp:142/:172 兩個實例）。分寸：陳述事實、不貶前人——無聲缺陷正常使用不會露出來。「11 個」是第一輪驗證的數字（CHANGELOG 原文 eleven defects）；後續審查坐實的追加條目以「其後又坐實多條」帶過，不改 11——S 報告曾請 Adam 裁定是否改總數，初版採此表述。")

# 17 — Crash class
s = slide_base("b")
title(s, "Crash 類（4 條）")
y = 1.62
row(s, y, "1", "一個 POST 打死整個 kernel", "/ndt/intent_translator/text 對 null translator 解參考——baseline 不用 AI 時 translator 就是 nullptr 且 handler 無 guard；null deref 是 signal，下方的 catch 接不到。修成 guard 回 503。", h=1.16); y += 1.22
row(s, y, "2", "Ryu 的 table-miss 規則讓 Classifier crash", "\"actions\": [] 對空 vector 呼叫 .front()。陷阱：呼叫點在 SPDLOG_LOGGER_TRACE 裡——spdlog 就算 level 關閉也會先求值參數。", h=1.0); y += 1.06
row(s, y, "3", "sFlow parser 缺邊界檢查", "對外可達的輸入面：任何能對 UDP:6343 發包的東西。ASan 證實拿掉檢查可重現 heap-buffer-overflow。", h=1.0); y += 1.06
row(s, y, "4", "缺欄位的控制平面回應殺掉 kernel", ".value(\"dpid\",\"\") → stoull 丟 invalid_argument；唯一的 catch 只接 json::parse_error（三個 ingest 同款）。為什麼沒人發現：baseline 只在啟動抓一次拓撲，曝險就那一次；改成週期輪詢後，每次輪詢都在骰。", h=1.3)
notes(s, "template Page 15 ＋ A2 新增的 f21d7a0（④）。素材：fbd8140、f21d7a0 commit message；CHANGELOG 條目 2、3。④ 的誠實註記已寫進投影片（缺陷是既有的、曝光度是新的——與 ① 的 --no-ai、「併發與生命週期」頁 ② 的週期刷新同款處理）。baseline 驗證：28b8b13:TopologyAndFlowMonitor.cpp:313-314、catch :340/:421/:504。")

# 18 — Data correctness I
s = slide_base("b")
title(s, "資料正確性類（一）")
y = 1.62
row(s, y, "1", "flow table 資料競爭", "handlePacket 無鎖查表、find-then-branch 非原子——兩個 sFlow worker 撞同一新 key 時，輸家的賦值蓋掉贏家的 bytes。一個鎖現在涵蓋查找、兩個分支與其間的診斷。", h=1.16); y += 1.22
row(s, y, "2", "unsigned counter delta 下溢 → 1.8×10¹⁹ bps 直達 API", "counterDelta 現在飽和於零。elephant-flow 旗標的清除 else 被註解掉——掉一次 update 就永久鎖死，已還原。", h=1.0); y += 1.06
row(s, y, "3", "macToUint64 對壞 MAC 默默回錯值", "\"00:11:22:33:44:5\" → 73588229125，查到錯的 host。修成驗證後回 optional，每個呼叫端都繼承檢查。", h=1.0); y += 1.06
row(s, y, "4", "無法解析的 /stats/flow body 回 json::array()", "與「這台 switch 沒規則」無法區分——套用後會把該 dpid 的規則全部掃掉。現在回 nullopt。", h=1.0)
notes(s, "template Page 16（原四條）。素材：CHANGELOG L390-408（themed 條 1、2、4、6）。")

# 19 — Data correctness II
s = slide_base("b")
title(s, "資料正確性類（二）")
y = 1.75
row(s, y, "5", "丟掉 boost::edge() 的 found 旗標", "兩處逐字相同：取 .first、丟 .second——把 miss 包進一個 engaged 的 optional，對「找不到」回答「有」。只要圖上有單向鏈路，per-edge 流量記帳就跑在無效 descriptor 上。不是 race——sanitizer 看不到，也不留 log。", h=1.5); y += 1.66
row(s, y, "6", "VLAN 欄位名不符：一顆未引爆的地雷", "guard 查 vlan_vid、實際讀 vlan_id——後者在 OpenFlow、Ryu 輸出、整個 repo 都不存在。任何人安裝的第一條 VLAN 規則，會讓該 switch 在同一條 flow 上永遠死掉：mark-and-sweep 停擺、排在後面的 switch 全部不再更新。今天沒人發 VLAN match，正是它會以謎題而非回歸的形式出現的原因。已立案 issue #3——只記錄，不修。", h=1.9)
add_text(s, 0.6, 5.7, 12.1, 0.5, [dict(text="「修一個還沒有人踩到的 bug，並說明為什麼它一定會被踩到。」", size=13, italic=True, color=C["red"])])
notes(s, "template Page 16 續（A2 2026-08-13 新增：cbb504a、5054249，皆已對 28b8b13 原文驗證——:857/:879 的 .first、Classifier.cpp:771/:773 的 vlan_vid/vlan_id）。VLAN：Adam 裁定只記錄不修（只修一側比不修更糟）。⑥ 其實也是「只會加、不會刪」家族的近親（mark-and-sweep 停擺）。")

# 20 — Add-only family
s = slide_base("b")
title(s, "「只會加、不會刪」——同一個 bug 形狀重複出現")
kicker(s, "每個資料攝取點都該問同一個問題：舊資料什麼時候消失？", y=1.4)
y = 2.0
row(s, y, "1", "拓撲只在啟動抓一次（88 ms）就再也不讀", "修成 5s/30s 輪詢。", h=0.72); y += 0.78
row(s, y, "2", "Classifier 跳過空 flow table", "被移除的規則永不被掃掉。", h=0.72); y += 0.78
row(s, y, "3", "setAllPaths 從不清 map", "斷鏈已刪的路徑，還在回答 get_path_switch_count。", h=0.72); y += 0.78
row(s, y, "4", "Answer::from_json——已發現、未修", "誠實列出。", h=0.72, color=C["slate"]); y += 0.78
row(s, y, "5", "OVS 路由規則永不刪除——已查證、刻意未修", "既有 Ryu 控制程式零 OFPFC_DELETE、add_flow 不設 timeout；switch 重連也不重算。修法不明顯安全：單向故障時刪規則會把還能走的路徑變成 table miss——處置待裁決。", h=1.0, color=C["slate"])
add_text(s, 0.6, 6.15, 12.1, 0.45, [dict(text="共同危害：孿生用已不存在的狀態計算——比空答案更糟，因為它看起來很有自信。", size=13, bold=True, italic=True, color=C["red"])])
notes(s, "template Page 17 ＋ 第 5 列（2026-08-13 cfbbf24 查證的最新一例，template 校訂之後才坐實）。歸類照 2c81b26 先例「既有 Ryu 控制程式的缺陷」；要不要上台、以及三個處置選項（只記錄／重連時補重算／連 BFS 可達性一起處理）都等 Adam 裁決，初版先列出並標「未修」。詳見 doc/2026-07-29_HANDOFF.md §2b 第 8 列。素材：CHANGELOG L398-401、條目 28、29；71d27c1、820c2a2、eb9c860。「同一形狀找到多例」展示的是系統性歸納，不只是逐條修——值得口頭強調。")

# 21 — State destruction I
s = slide_base("b")
title(s, "狀態破壞與誤報類（一）")
y = 1.62
row(s, y, "1", "OVS liveness 把整組機組報死", "ovs-vsctl list-br「失敗」與「回報零 bridge」不可區分——都是空 vector，pclose 狀態沒人看：一次掉線、十台全標死。且該分支只呼叫 setVertexDown 從不 setVertexUp——「死」是永久的。一次事故 3596 行 log。", h=1.4); y += 1.46
row(s, y, "2", "powerOff 先毀掉 powerOn 要用的紀錄", "先用讀不到的（空的）port list 蓋掉 graph 存檔、再刪 bridge——兩份要接回的紀錄都毀了。baseline 的 listPorts 對 popen 失敗回空、pclose 狀態丟棄。修成 std::optional：port list 未知，就拒絕整個操作。", h=1.4); y += 1.46
row(s, y, "3", "synthetic power 報 1.9×10¹⁴ 瓦", "uniform_int_distribution<uint64_t>(0, 2⁶⁰) 每次 poll 重擲；還有兩份獨立 RNG，讓兩個 API 對同一台 switch 同時給不同答案。修成種子化的 30–150 W 穩定值。", h=1.4)
notes(s, "template Page 18 ①②③。素材：CHANGELOG 條目 14、15、25；themed 條 7。② 已對 28b8b13 原文驗證。")

# 22 — State destruction II
s = slide_base("b")
title(s, "狀態破壞與誤報類（二）——電源回報家族")
y = 1.62
row(s, y, "4", "TESTBED 電源：同一個操作、兩份實作，兩份都說謊", "可達的那份把 curl 的 exit code 當結果——gateway 回 500 / 401 / 錯誤頁都回報 Success，而且從不更新 graph；另一份會照 caller 要求改 graph、不看 gateway 的答案——但它從寫下那天到刪掉那天都是零呼叫點。修復：判定抽成純函式 interpretRelayResponse，放進真正會跑的路徑。", h=1.62); y += 1.68
row(s, y, "5", "十台 switch 共用同一個插座", "兩個分支寫死同一個 smart_plug_ip（實為 s2 的）；拓撲檔一直帶著逐台真實指派，載入器直接讀過去。在真硬體上，信任這個端點的消費端會對十台 power-cycle 同一個錯的插座。沒被發現的原因：kernel 自己的電源路徑不用這些欄位——沒有東西會牴觸這個虛構。", h=1.42); y += 1.48
row(s, y, "6", "電源指令回報 ok、不看結果", "IntentTranslator 把 setSwitchPowerState 當裸述句呼叫、丟掉回傳的 bool；拓撲解析不出的裝置名，整個跳過呼叫——兩者都掉進共用的 return \"ok\"。", h=1.1)
notes(s, "template Page 18 ④ ＋ A2 新增 59dc5d3、916d330 的 baseline 半邊。⚠ ④ 的寫法是本次草稿的裁定候選：template 原 ④（「一律把 graph 改成 caller 要求的狀態」）與 3292653 message（「從未碰過 graph」）的矛盾，已用 baseline 原文解開——兩句都對，指不同 overload：28b8b13 的 :368 三參數版（scrape HTML 後照 action 改 graph、scrape 結果只寫 log）是死碼、零呼叫點；:689→:716 可達版是裸 curl rc==0、不碰 graph。初版合寫成一條完整故事；A2 原本把 3292653 列排除（理由即該矛盾），此寫法待 Adam 確認。59dc5d3 baseline 驗證：:1958/:1972 兩分支寫死同一 IP（投影片上不放 IP）。916d330 的 OpResult 半邊照 8c25dbc 先例在功能頁。好的串場句：這一層就坐在 Phase 7 那條鏈的正上方——helper 沒安裝的 P4 powerOff，對使用者回答「ok」。")

# 23 — Unidirectional blackhole
s = slide_base("b")
title(s, "本節最強的一條：單向鏈路故障，讓路由重算永久崩潰", size=24)
bullets(s, [
    dict(t="場景：只斷一個方向——控制面看見不對稱，而配對事件永遠不會來：不對稱是穩態，不是暫態", size=13.5),
    dict(t="機制：既有 Ryu 控制程式的 DiGraph 進入不對稱穩態 → BFS 撞 KeyError → 路由重算從此永久崩潰、流量無限期黑洞", size=13.5),
    dict(t="修法：BFS 只走雙向都存在的邊——繞路自動發生，且嚴格不弱於舊行為（舊 crash 案變繞路、舊不可達案照舊）", size=13.5),
], y=1.72, h=2.1)
stat(s, 0.8, 4.0, 3.9, 1.9, "291 秒", "100% 丟包、零自癒（修復前，OVS 路實測）", num_color=C["red"], num_size=32)
stat(s, 5.0, 4.0, 3.9, 1.9, "9–15 Mbps", "黑洞期間 twin 同時回報該 flow 的流量；edges_up 287/288", num_color=C["red"], num_size=26)
stat(s, 9.2, 4.0, 3.4, 1.9, "12.5 秒", "同款故障，P4 路自癒（第 5 節對照）", num_color=C["teal"], num_size=32)
add_text(s, 0.6, 6.2, 12.1, 0.4, [dict(text="缺陷在實驗室既有的 Ryu 控制程式（intelligent_router.py，隨初始 commit 帶進 repo）", size=11, color=C["faint"])])
notes(s, "template A2 新增裁定表首列（034da18）＋ Adam 裁決「單向斷鏈進 bug 頁」。素材：034da18 完整 message；實測數字出自 C-live-ovs-runbook（291 秒）與 W2-live-reverify（P4 側 12.5 秒）。twin 誤報細節：黑洞期間 twin 報該 flow 9–15 Mbps 流動、edges_up 287/288。這條與 Failover 頁的 2c81b26（整支沒有 remove_edge）同源，都註明「既有 Ryu 控制程式的缺陷」。")

# 24 — API errors & resources
s = slide_base("b")
title(s, "API 錯誤處理與資源類（5 條）")
y = 1.56
row(s, y, "1", "客戶端錯誤回 500 家族", "?src_ip=not.an.ip、?dpid=abc 一路丟到最外層 catch；{not json 回 202「已接受」；hex eth_type（OpenFlow 工具的正常寫法）被 400 拒。修復含 tryParseUint64——stoull 會把 \"12abc\" 讀成 12、\"-1\" 繞成 2⁶⁴−1。", h=1.06, body_size=11); y += 1.1
row(s, y, "2", "log 洪水淹掉真正的那一行", "path-walk 迴圈每毫秒警告一次：一個設錯的 port 產出 270,991 行、41 MB。修成 KeyedFailureLog 邊緣觸發回報。", h=0.82, body_size=11); y += 0.86
row(s, y, "3", "閒置 100% CPU", "sFlow 迴圈 poll() 用 0 ms timeout 忙轉。", h=0.62, body_size=11); y += 0.66
row(s, y, "4", "無界的 curl 請求（與穩健性頁是同一事件的兩半）", "kernel 讀每台 switch flow table 的 curl -s，從 baseline 起就沒有 --max-time，而 sweep 逐台序列——一台不回應，連帶卡住後面每一台。實測（SIGSTOP 一台 bmv2）：請求永不返回。為什麼沒人發現：Ryu/OVS 從不「活著但不回應」——是 P4/bmv2 帶進新故障模式，引爆躺了很久的缺陷。修成 --max-time 8，三個對外請求抽成可斷言的 builder（7 條測試）。", h=1.5, body_size=11); y += 1.54
row(s, y, "5", "API key 被寫進 log", "產品的 Intent Translator 啟動即把 OPENAI_API_KEY 以 INFO 寫進 log；且那行跑在 null 檢查之前——沒設 key 時先把 null 餵進 fmt，本該解釋設定錯誤的 ERROR 到不了。", h=0.95, body_size=11)
notes(s, "template Page 19（含 A2 裁定的 1a7d815 kernel 半邊＝④，與穩健性頁互為一體兩面，兩頁都有明說）＋ A2 新增 57346f6（⑤）。⑤ 的 A1.3 辨析（重要）：這是產品自己的 Intent Translator 功能在用 OpenAI——不是被禁止的「開發流程」話題，不要砍。baseline 驗證：LLMAgent.cpp:27-32（log 在 null 檢查之前）；DeviceConfigurationAndPowerManager baseline :507 無 --max-time。素材：CHANGELOG 條目 22、23、24、31；themed 條 5；L382；1a7d815、57346f6。")

# 25 — Concurrency & lifecycle
s = slide_base("b")
title(s, "併發與生命週期類（2 條）")
y = 1.75
row(s, y, "1", "FlowDispatcher 三個生命週期缺陷（原本零測試）", "stop() 無鎖迭代並清空 workers，而 enqueue() 持鎖插入；running_ 在鎖外寫入——lost wakeup、stop() 自身死鎖；stop() 搬走 map 後 enqueue() 還會生 worker，留下無主的 joinable thread。補 6 個生命週期測試。", h=1.6); y += 1.75
row(s, y, "2", "宣告了、但從未使用的 mutex", "m_allPathMap / m_switchCountMap 六處存取無同步。誠實帶一句：baseline 是啟動期 race；我加的週期性刷新讓它變成永久 race——缺陷是既有的，曝光度是新的。", h=1.4)
notes(s, "template Page 20。素材：CHANGELOG 條目 18、20（d5f5bfa、0596dd1）。44fa86e（m_ipStrToDpidMap operator[] 讀取即插入→dpid 0，已驗 baseline 九處）依 S §5.9 暫不進頁面——要不要擴充由 Adam 決定。")

# 26 — Ryu wedge story
s = slide_base("b")
title(s, "調查故事：Ryu flow-stats wedge")
bullets(s, [
    dict(t="症狀：Mininet 活著時重啟 Ryu → /stats/flow 從此永遠回空表", size=13.5),
    dict(t="特徵化：兩次重現、151 個樣本（doc/audit/2026-08-07_ryu-wedge-trace.tsv）；四個假說全數證偽——root cause 至今未證明，照實講", size=13.5, b=True),
    dict(t="不碰 Ryu、修掉危害：wedged 回覆要 1.011 秒（健康 0.027–0.083 秒；1.0 秒正是 ryu/lib/ofctl_utils.py 的 DEFAULT_TIMEOUT）——所以 ≥0.5 秒的空表一律拒收、沿用前值", size=13.5),
    dict(t="操作守則：不要單獨重啟 Ryu——要重啟就連 Mininet 一起", size=13.5),
    dict(t="這頁展示的方法：假說 → 證偽 → 承認未知 → 只修危害", size=13.5, c=C["teal"]),
])
notes(s, "template Page 21（可略，但對教授場合是加分頁）。素材：CHANGELOG L412-419；tsv 檔。")

# 27 — Tech stack
s = slide_base("t")
title(s, "技術棧總覽")
stacks = [
    ("Kernel（C++23）", ["CMake", "Boost.Beast / Asio / URL（HTTP）", "nlohmann::json · spdlog · libssh", "GTest", "ASan / TSan sanitizer 建置"], C["steel"], C["steelTint"]),
    ("P4 資料面", ["P4_16 / v1model", "bmv2 simple_switch_grpc", "p4c-bm2-ss", "P4Runtime（gRPC / protobuf）", "PRE clone session · Mininet"], C["teal"], C["tealTint"]),
    ("Proxy 與控制面（Python）", ["FastAPI / uvicorn", "grpc", "Ryu / OpenFlow 1.3", "sFlow v5 · LLDP", "SNMP"], C["slate"], C["slateTint"]),
]
for i, (hh, items, col, tint) in enumerate(stacks):
    x = 0.7 + i * 4.15
    add_box(s, x, 1.85, 3.9, 4.3, fill=tint, radius=0.1)
    add_text(s, x + 0.28, 2.1, 3.35, 0.5, [dict(text=hh, size=15, bold=True, color=col)])
    add_text(s, x + 0.28, 2.75, 3.4, 3.2,
             [dict(text=t, bullet=True, size=12.5, gap=8) for t in items])
notes(s, "template Page 22。素材：CMakeLists.txt、p4_proxy/requirements.txt、各 SPEC.md。依 A1 規則，不出現 AI 協作相關字樣。")

# 28 — Methodology
s = slide_base("t")
title(s, "方法論——本次工作真正的重點技術")
y = 1.7
row(s, y, "1", "差異測試", "OVS 路是已知良好的——直接拿它的行為當 P4 的規格（L4 differential）。", color=C["teal"], h=0.95); y += 1.08
row(s, y, "2", "Mutation testing 作為驗收關卡", "每個測試都要親眼看過它失敗一次才算數：弄壞實作 → 看它紅 → 修回，證據入檔 doc/audit/。這道關卡抓到 11+ 個「通過、但什麼都沒證明」的測試。", color=C["teal"], h=1.15); y += 1.28
row(s, y, "3", "allowlist 閘門", "沒列入白名單的 warning／差異一律 fail——新問題藏不進舊雜訊。", color=C["teal"], h=0.95); y += 1.08
row(s, y, "4", "證據式設計", "liveness 三態、南向 OpResult——「無法判斷」與「失敗」，都不准偽裝成「成功」。", color=C["teal"], h=0.95)
notes(s, "template Page 23。素材：doc/2026-08-07_testing_tools_overview.md 方法論段；CHANGELOG L433-441。「11+」出處：CHANGELOG L438 記 11（2026-08-08 時點）；其後輪次未逐條重數，維持下限表述——A3 重數規則對此條的極限，已如實標注。S 報告另建議加第 ⑤ 點「引用會腐爛」（行號改 symbol、烘進文件的數字移除；佐證：某 runbook 全檔唯一存活的引用就是唯一用 symbol 的那處）——template 校訂未採，要不要加由 Adam 決定。")

# 29 — Five layers
s = slide_base("q")
title(s, "五層測試架構（L0–L4）")
layers = [
    ("L0", "建置檢查", "跨 repo 破壞，分鐘級發現"),
    ("L1", "單元測試", "ctest 與直接執行兩種方式都跑——各自會漏掉對方能抓的失敗（ctest 每 case 獨立 process，SetUpTestSuite 失敗會被吃掉）"),
    ("L2", "API 契約", "結構＋語意不變量＋錯誤路徑；期望值從拓撲檔推導，不是寫死"),
    ("L3", "元件契約", "改一個 endpoint，能回答「誰會壞」——blast radius"),
    ("L4", "OVS/P4 差異比對", "allowlist 三分類：允許的 P4 限制／數值容忍／其餘就是 bug"),
]
y = 1.78
for l, hh, bb in layers:
    add_box(s, 0.7, y, 1.0, 0.78, fill=C["slate"], radius=0.08)
    add_text(s, 0.7, y, 1.0, 0.78, [dict(text=l, size=16, bold=True, color=C["white"], align=PP_ALIGN.CENTER)], valign=MSO_ANCHOR.MIDDLE)
    add_text(s, 1.95, y - 0.02, 10.6, 0.9, [
        dict(text=hh, size=13.5, bold=True, gap=2),
        dict(text=bb, size=11.5, color=C["gray"], ls=1.08),
    ], valign=MSO_ANCHOR.MIDDLE)
    y += 0.92
add_text(s, 0.7, 6.45, 11.9, 0.4, [dict(text="stack.sh 依模式（ovs / p4）用正確順序編排啟動——測試環境本身也是被管理的", size=11.5, color=C["faint"])])
notes(s, "template Page 24。素材：doc/2026-08-07_testing_tools_overview.md（首選）、tools/test_workflow/README.md、tools/contract_test/README.md。")

# 30 — Test assets
s = slide_base("q")
title(s, "測試資產規模與品質保證")
stat(s, 0.7, 1.75, 3.9, 1.7, "554", "gtest cases · 44 個 .cpp（ctest 與直跑一致）", num_color=C["slate"])
stat(s, 4.85, 1.75, 3.9, 1.7, "634", "Python 測試（proxy 402 ＋ kernel 契約 232）", num_color=C["slate"])
stat(s, 9.0, 1.75, 3.6, 1.7, "6", "shell 測試（含故障注入 harness）", num_color=C["slate"])
bullets(s, [
    dict(t="品質保證是機制、不是數字：每個測試都有 mutation 證據（弄壞過、紅過、修回）", b=True, size=13),
    dict(t="修掉 4 個測試工具自己的假 PASS——例：unittest 把 skipped 算進 Ran N，整檔 skip 也顯示綠燈；另修 3 個從未真正執行的測試", size=13),
    dict(t="log 判定：allowlist ＋ FORBID ＋ 崩潰偵測——crash 訊息不吃 allowlist", size=13),
    dict(t="成長不是灌水：426 → 554 的主要跳升來自一輪密集審查與 Phase 7——每個修復都帶著自己的迴歸測試進來", size=13),
], y=3.85, h=2.9)
notes(s, "template Page 25。數字 2026-08-13 於 head cfbbf24 依 A3.1 指令重量：gtest grep 554（TEST_P=0 已驗，grep 數＝runtime 數；實跑 554/70 suites 一致）、p4_proxy/tests 402（15 檔）、tests/python 232（7 檔，未註冊 ctest、獨立執行）、shell 6。成因（教授若問 426→554 怎麼來的）：密集審查輪為三個實測缺陷各補迴歸測試＋四組新工具（fuzz harness、P4 覆蓋閘門、故障注入、twin 對帳）各帶測試＋test_RequestDeadlines 7 條。素材：CHANGELOG L433-441；doc/audit/ mutation-evidence 系列；fbd8140。")

# 31 — Documentation assets
s = slide_base("q")
title(s, "文件資產——接手的人有路可走")
cats = [
    ("操作", ["full_test_runbook", "ovs_manual_test_runbook", "p4_manual_test_runbook", "environment_gotchas"]),
    ("設計／狀態", ["p4_bmv2_support_plan（分 phase）", "phase7_power_mechanism_design", "p4_status_and_test_guide", "HANDOFF · ndt_api"]),
    ("測試", ["testing_tools_overview", "testing_workflow", "test_coverage_gaps（誠實列出還沒測的）"]),
    ("調查紀錄", ["doc/audit/：wedge trace", "mutation 證據", "獨立審查與裁決紀錄"]),
]
for i, (hh, items) in enumerate(cats):
    x = 0.7 + (i % 2) * 6.25
    y = 1.85 + (i // 2) * 2.35
    add_box(s, x, y, 5.95, 2.1, fill=C["card"], line_color=C["line"], radius=0.09)
    add_text(s, x + 0.25, y + 0.12, 5.4, 0.4, [dict(text=hh, size=13.5, bold=True, color=C["slate"])])
    add_text(s, x + 0.25, y + 0.55, 5.5, 1.45,
             [dict(text=t, bullet=True, size=11, gap=4) for t in items])
notes(s, "template Page 26（可略，一頁帶過）。檔名為現行名（doc/ 已於 2026-08-13 全面改為「建立日期_原名」，投影片省略日期前綴求可讀；口頭可提命名紀律：檔名帶建立日、ls 即時序）。")

# 32 — L4 differential
s = slide_base("l")
title(s, "L4 差異比對：P4 對齊 OVS baseline")
add_text(s, 0.6, 1.8, 12.1, 1.6, [
    dict(text="結論：PASS", size=20, bold=True, color=C["teal"], gap=6),
    dict(text="P4 與 OVS baseline 一致；14 個被接受的差異全部有文字記錄；12 條因 P4 補齊而過時的 allowlist，逐條對活系統驗證後移除。", size=13.5, ls=1.15),
])
placeholder(s, 0.6, 3.6, 12.1, 2.9, "比對輸出截圖、14 個被接受差異的分類表。素材指標：commit dac192b；CHANGELOG 條目 17；tools/contract_test/baseline_diff_allowlist.txt")
notes(s, "template Page 27（骨架預填，A1.4）。")

# 33 — Failover E2E
s = slide_base("l")
title(s, "Failover 端到端驗證——同一個故障，兩個資料面")
add_text(s, 0.6, 1.5, 12.1, 0.42, [dict(text="注入只斷單一 link 的故障（tc netem loss 100% 雙端），流量改道、恢復後路由回原路，全程孿生狀態正確。", size=12.5)])
tbl_data = [
    ("量測", "bmv2（P4）", "OVS（對照組）"),
    ("偵測到 link down", "2 筆通知、零假訊", "+42.6 秒、2 筆零假訊"),
    ("ping 中斷後自癒", "16.63 秒", "52.42 秒"),
    ("復原（鏈路回來）", "hitless、0 遺失", "hitless"),
    ("繞路後路徑", "1 → 6 → 10 → 7 → 4（規則層讀回確認 OUTPUT:3→4）", "—"),
    ("單向故障（只斷一個方向）", "12.5 秒自癒、規則改 OUTPUT:1→2", "291 秒 100% 丟包、零自癒（修復前）"),
]
gf = s.shapes.add_table(len(tbl_data), 3, Inches(0.6), Inches(2.05), Inches(12.1), Inches(0.46 * len(tbl_data)))
tbl = gf.table
# plain grid style
tbl_el = gf._element.graphic.graphicData.tbl
style_ids = tbl_el.findall(qn("a:tblPr"))
if style_ids:
    for sid in style_ids[0].findall(qn("a:tableStyleId")):
        style_ids[0].remove(sid)
    sid_el = style_ids[0].makeelement(qn("a:tableStyleId"), {})
    sid_el.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    style_ids[0].append(sid_el)
tbl.columns[0].width = Inches(3.3)
tbl.columns[1].width = Inches(4.6)
tbl.columns[2].width = Inches(4.2)
green_cells = {(1, 1), (2, 1), (3, 1), (5, 1)}
red_cells = {(5, 2)}
for ri, rdata in enumerate(tbl_data):
    tbl.rows[ri].height = Inches(0.46)
    for ci, val in enumerate(rdata):
        cell = tbl.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(C["navy"] if ri == 0 else (C["card"] if ri % 2 == 0 else C["white"]))
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tfc = cell.text_frame
        tfc.word_wrap = True
        pc = tfc.paragraphs[0]
        rc = pc.add_run()
        rc.text = val
        if ri == 0:
            set_run(rc, size=12, bold=True, color=C["white"])
        elif (ri, ci) in green_cells:
            set_run(rc, size=11.5, bold=True, color=C["teal"])
        elif (ri, ci) in red_cells:
            set_run(rc, size=11.5, bold=True, color=C["red"])
        else:
            set_run(rc, size=11.5, color=C["ink"])
add_text(s, 0.6, 5.15, 12.1, 0.55, [dict(text="這張表本身就是敘事：同一個故障形狀，兩個資料面的差異來自實作、而非設定——右下那格正是第 2 節最強的那條 bug。", size=12, bold=True, italic=True, color=C["navy"])])
placeholder(s, 0.6, 5.85, 12.1, 1.0, "斷鏈前後的路徑圖（數據已備齊，缺視覺化）。素材指標：ca72d22；doc/2026-08-10_p4_manual_test_runbook.md")
notes(s, "template Page 28（數據已有：A-live-runbook、C-live-ovs-runbook、W2-live-reverify，repo 外 ~/Documents/NDTwin documentation/）。斷鏈手法：s6-eth3↔s9-eth2 tc netem loss 100% 雙端；不用 ifconfig down——它會弄壞整台 bmv2。291 秒那格＝「單向鏈路故障」頁的缺陷（修復前 OVS 路），修復後行為見 034da18 迴歸測試——口頭要講清楚時序。")

# 34 — 200 Mbps
s = slide_base("l")
title(s, "200 Mbps 壓力測試＋負載下斷鏈")
add_text(s, 0.6, 1.8, 12.1, 0.5, [dict(text="一句話結論：OVS stack 推到 200 Mbps 並在負載下斷鏈，記錄哪些量測撐住、哪些開始失真。", size=14)])
placeholder(s, 0.6, 2.6, 12.1, 3.6, "吞吐／量測誤差圖表。素材指標：commit cc7437a 完整 message。")
notes(s, "template Page 29（骨架預填）。")

# 35 — Flow-rate accuracy
s = slide_base("l")
title(s, "孿生 flow-rate 精確度量測")
add_text(s, 0.6, 1.8, 12.1, 0.5, [dict(text="一句話結論：量化孿生回報的 flow rate 與實際流量的誤差，以及誤差開始失效的邊界條件。", size=14)])
placeholder(s, 0.6, 2.6, 12.1, 3.6, "誤差曲線、邊界條件表。素材指標：commit d7bf52f 完整 message；doc/audit/ 相關 runbook。")
notes(s, "template Page 30（骨架預填）。理論地板參考（口頭備用）：sFlow 1/256 取樣誤差 196√(1/c)。")

# 36 — Independent cross-check
s = slide_base("l")
title(s, "獨立交叉檢查")
bullets(s, [
    dict(t="由不了解實作的第三方視角，對活著的 OVS stack 做 51 次唯讀檢查", size=14),
    dict(t="12 項發現全數裁定——1 項真 bug（已修），其餘為已記錄行為或誤報，逐條查證後結案", size=14),
    dict(t="另一輪 47 項 HIGH 的審查裁定出 21 項真問題——全部逐條查證、逐條有下落", size=14),
], y=1.8, h=2.4)
placeholder(s, 0.6, 4.4, 12.1, 1.9, "是否深入展開由 Adam 決定，預設一頁帶過。素材指標：commit 2bec2a5、fbd8140。")
notes(s, "template Page 31。依 A1.3：敘述用「第三方交叉檢查／獨立驗證流程」，不提執行工具。後續還有一輪 55 條、實質全中零誤判的審查（Tier 1 四條全修）——Adam 已裁決不另開頁、發現併進既有頁、來源不提，故本頁不列；其產出的 baseline 條目已散在第 2 節。")

# 37 — Demo
s = slide_base("l")
title(s, "Demo")
placeholder(s, 0.6, 1.7, 12.1, 1.35, "demo 腳本與畫面。")
add_text(s, 0.6, 3.3, 12.1, 0.45, [dict(text="三個候選場景（皆已在 live 環境驗過可行）：", size=13.5, bold=True)])
y = 3.85
row(s, y, "1", "斷鏈 failover 即時展示", "16.63 秒自癒 ＋ hitless 復原，數字現成。", color=C["navy"], h=0.8); y += 0.88
row(s, y, "2", "Web-GUI 看 P4 模式拓撲與流量", "上層零修改的最直觀證明——GUI 根本不知道下面換成了 bmv2。", color=C["navy"], h=0.8); y += 0.88
row(s, y, "3", "killed switch 的 liveness 三態變化", "從清單消失、10 Hz 取樣零 up-blip。", color=C["navy"], h=0.8)
notes(s, "template Page 32。Demo 腳本目前不存在（全 repo find 過零命中）——腳本與畫面待 Adam。三候選皆出自 runbook 可直接演。")

# 38 — Summary (dark)
s = slide_base("l", dark=True)
title(s, "總結與未完成事項", dark=True)
add_text(s, 0.7, 1.55, 11.9, 1.15, [
    dict(text="成果一句話：", size=14, bold=True, color=C["iceOnDark"], gap=3),
    dict(text="P4/bmv2 資料面完整接入、上層零修改；baseline 共用路徑的無聲缺陷修復；五層測試架構讓兩個資料面都有可判定的健康標準。", size=15, bold=True, color=C["white"], ls=1.15),
])
add_text(s, 0.7, 2.9, 6, 0.4, [dict(text="誠實列出未完成：", size=13, bold=True, color=C["iceOnDark"])])
unfinished = [
    "Phase 3 未開始（其餘 Phase 完成；Phase 5 counter-sample 半邊未做——與 OVS 行為對等）",
    "Answer::from_json 第四例、OVS 路由 stale rule——已發現，處置待定",
    "Ryu wedge root cause 未證明；hopsCounter 分母膨脹已記錄未修",
    "VLAN 欄位名不符已立案（issue #3）未修",
    "已知風險：southbound 指令仍以 popen(\"curl …\") 拼接、單引號不跳脫——刻意延後，22 處／3 檔",
    "bmv2 本身接受未取得 mastership 的 pipeline push——實測坐實；我們這側已有防護，尚未回報上游",
]
add_text(s, 0.7, 3.35, 11.9, 2.6,
         [dict(text=t, bullet=True, size=12, color=C["iceOnDark"], gap=6, ls=1.1) for t in unfinished])
add_text(s, 0.7, 6.15, 11.9, 0.5, [dict(text="下一步：合併實驗室上游的 sharding 變更（已裁定延至報告後）· 長時 soak／drift 量測 · demo 定稿", size=12.5, bold=True, color=C["white"])])
notes(s, "template Page 33（未完成清單按 2026-08-13 現況核對：Phase 7、8 都已完成——template 校訂前的舊版寫錯方向，這裡已修正）。「誠實列出未完成」在教授場合是加分不是扣分。OVS stale rule 為 cfbbf24 新查證項，template 未及收錄，已按其精神列入。mastership 上報上游前還缺第三方 client 重現（p4runtime_lib 的 MasterArbitrationUpdate 會永遠阻塞，需手寫低階 gRPC）。")

out = sys.argv[1] if len(sys.argv) > 1 else "NDTwin-progress-report-draft-v1.pptx"
prs.save(out)
print("written:", out, "slides:", page_no)
